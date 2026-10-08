"""
========================================================
  Dental AI Tool – Data Manager Module
  Quản lý dữ liệu, sao lưu tự động, xuất Excel & JSON
========================================================
"""

import os
import csv
import json
import re
import shutil
import zipfile
import io
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Cấu hình đường dẫn ──────────────────────────────────
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_CSV = os.path.join(BASE_DIR, "ai_evaluation_results.csv")
OUTPUT_EXCEL = os.path.join(BASE_DIR, "ai_evaluation_results.xlsx")
PROGRESS_JSON = os.path.join(BASE_DIR, "session_progress.json")
RESPONSES_DIR = os.path.join(BASE_DIR, "ai_responses")
BACKUPS_DIR = os.path.join(BASE_DIR, "backups")

# 16 cột chuẩn của nghiên cứu
CSV_HEADERS = [
    "Case_ID", "AI_Model", "Prompt_Type", "Time_Seconds",
    "Pell_Gregory_Class", "Pell_Gregory_Position", "Pell_Gregory_Full",
    "Winter_Class", "Confidence", "Pederson_Level",
    "Hallucination_Flag", "Reasoning_Quality", "Operator",
    "Timestamp", "Session_Notes", "AI_Raw_Response"
]

# 4 người nhập số liệu chính thức
OPERATORS_LIST = [
    "GIANG",
    "HOÀNG",
    "MINH",
    "HÂN"
]

# 3 mô hình AI đánh giá
MODELS_LIST = [
    "GEMINI",
    "CLAUDE",
    "CHATGPT"
]

# Tổng số ca nghiên cứu (129 ca: Case_001 -> Case_129)
TOTAL_CASES = 129
ALL_CASES = [f"Case_{str(i).zfill(3)}" for i in range(1, TOTAL_CASES + 1)]

PG_CLASSES = ["I", "II", "III", "Không xác định"]
PG_POSITIONS = ["A", "B", "C", "Không xác định"]

WINTER_VALUES = [
    "Mesioangular", "Vertical", "Horizontal",
    "Distoangular", "Buccolingual", "Others"
]

PEDERSON_LEVELS = ["Nhẹ (3–4)", "Trung bình (5–6)", "Khó (7–10)"]

REASONING_LEVELS = [
    "0 – Không có lý luận",
    "1 – Lý luận mơ hồ",
    "2 – Có lý luận cơ sở",
    "3 – Lý luận chuẩn xác"
]

HALLUCINATION_LEVELS = ["Không", "Có – Nhẹ", "Có – Rõ ràng"]

PROMPTS = {
    "P1 – Basic": (
        "Please classify this impacted mandibular third molar based on the panoramic "
        "X-ray using Pell & Gregory and Winter classification systems. "
        "Also state your confidence level (0–100%)."
    ),
    "P2 – Structured Medical": (
        "You are an oral and maxillofacial surgery specialist. Analyze the provided panoramic "
        "radiograph step by step:\n\n"
        "Step 1 – Identify anatomical landmarks: Locate the anterior border of the ramus, "
        "the occlusal plane of the second molar, and the alveolar crest.\n\n"
        "Step 2 – Pell & Gregory Classification:\n"
        "  Class I: Crown fully anterior to ramus\n"
        "  Class II: Crown partially under ramus\n"
        "  Class III: Crown fully within ramus\n"
        "  Position A: Crown at/above occlusal plane of M2\n"
        "  Position B: Crown between occlusal plane and cervical line of M2\n"
        "  Position C: Crown below cervical line of M2\n\n"
        "Step 3 – Winter Classification: Determine angulation "
        "(Mesioangular / Vertical / Horizontal / Distoangular / Buccolingual / Others)\n\n"
        "Step 4 – Confidence: State your confidence level (0–100%).\n\n"
        'Output format (JSON):\n{\n  "P&G_class": "",\n  "P&G_position": "",\n'
        '  "Winter": "",\n  "confidence": 0,\n  "reasoning": ""\n}'
    )
}

PROMPT_TAGS = {
    "P1 – Basic": "P1: Prompt cơ bản kèm yêu cầu độ tin cậy – mô phỏng người dùng thông thường",
    "P2 – Structured Medical": "P2: Cấu trúc y khoa đầy đủ + Chain-of-Thought + JSON output"
}


def ensure_directories():
    """Đảm bảo các thư mục lưu trữ cần thiết tồn tại."""
    os.makedirs(RESPONSES_DIR, exist_ok=True)
    os.makedirs(BACKUPS_DIR, exist_ok=True)
    if not os.path.exists(OUTPUT_CSV):
        with open(OUTPUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerow(CSV_HEADERS)


def natural_sort_key(s: str):
    """
    Hàm tách số tự nhiên để sắp xếp đúng thứ tự (VD: Case_1, Case_2, Case_10).
    """
    match = re.search(r'\d+', str(s))
    return (int(match.group()), str(s)) if match else (999999, str(s))


def load_dataset(sort_by_case: bool = True) -> pd.DataFrame:
    """
    Đọc dữ liệu từ file CSV chính, chuẩn hóa cột theo CSV_HEADERS
    và sắp xếp theo thứ tự tự nhiên của Mã ca (Case_ID).
    """
    ensure_directories()
    if not os.path.exists(OUTPUT_CSV) or os.path.getsize(OUTPUT_CSV) == 0:
        return pd.DataFrame(columns=CSV_HEADERS)

    try:
        # Đọc linh hoạt bỏ qua số cột bất thường nếu có
        df = pd.read_csv(OUTPUT_CSV, dtype=str, on_bad_lines="skip")
        
        # Đảm bảo có đủ 16 cột
        for col in CSV_HEADERS:
            if col not in df.columns:
                df[col] = ""
        
        # Chỉ giữ đúng 16 cột chuẩn theo thứ tự
        df = df[CSV_HEADERS].fillna("")

        # Chuyển đổi kiểu dữ liệu cơ bản cho bảng tính
        df["Time_Seconds"] = pd.to_numeric(df["Time_Seconds"], errors="coerce").fillna(0).astype(int)
        df["Confidence"] = pd.to_numeric(df["Confidence"], errors="coerce").fillna(50).astype(int)

        # Sắp xếp theo thứ tự tự nhiên của Case_ID
        if sort_by_case and not df.empty and "Case_ID" in df.columns:
            df["_sort_key"] = df["Case_ID"].apply(lambda x: natural_sort_key(x)[0])
            df = df.sort_values(by="_sort_key").drop(columns=["_sort_key"]).reset_index(drop=True)

        return df
    except Exception as e:
        print(f"Lỗi khi đọc file CSV: {e}")
        return pd.DataFrame(columns=CSV_HEADERS)


def clear_dataset(create_backup: bool = True) -> bool:
    """
    Xóa toàn bộ dữ liệu bảng tính (giữ lại header chuẩn)
    và tự động tạo 1 bản snapshot dự phòng an toàn trước khi xóa.
    """
    ensure_directories()
    if create_backup and os.path.exists(OUTPUT_CSV) and os.path.getsize(OUTPUT_CSV) > 0:
        create_snapshot_backup(note="before_clear_all")

    with open(OUTPUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerow(CSV_HEADERS)
    return True


def delete_cases(case_ids: list, create_backup: bool = True) -> bool:
    """
    Xóa một hoặc nhiều ca bệnh cụ thể theo Case_ID khỏi cơ sở dữ liệu.
    """
    ensure_directories()
    df = load_dataset(sort_by_case=False)
    if not df.empty and case_ids:
        filtered_df = df[~df["Case_ID"].isin(case_ids)]
        save_dataset(filtered_df, create_backup=create_backup)
        return True
    return False


def get_case_record(case_id: str) -> dict:
    """
    Lấy thông tin bản ghi đánh giá mới nhất của một Case_ID cụ thể.
    """
    df = load_dataset(sort_by_case=False)
    if not df.empty and "Case_ID" in df.columns:
        matches = df[df["Case_ID"] == case_id]
        if not matches.empty:
            return matches.iloc[-1].to_dict()
    return {}


def upsert_evaluation_record(record: dict) -> bool:
    """
    Thêm mới hoặc cập nhật bản ghi đánh giá ca bệnh nếu Case_ID đã tồn tại.
    """
    ensure_directories()
    case_id = record.get("Case_ID", "")
    df = load_dataset(sort_by_case=False)

    if not df.empty and case_id in df["Case_ID"].values:
        # Cập nhật bản ghi đã có
        idx = df[df["Case_ID"] == case_id].index[-1]
        for col in CSV_HEADERS:
            if col in record:
                df.at[idx, col] = record[col]
        save_dataset(df, create_backup=True)
    else:
        # Thêm mới
        append_evaluation_record(record)

    # Lưu thêm file JSON metadata
    safe_model = re.sub(r'[^a-zA-Z0-9_-]', '_', str(record.get("AI_Model", "model")))
    safe_prompt = re.sub(r'[^a-zA-Z0-9_-]', '_', str(record.get("Prompt_Type", "prompt")))
    ts_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_filename = f"{case_id}_{safe_model}_{safe_prompt}_{ts_str}.json"
    json_path = os.path.join(RESPONSES_DIR, json_filename)
    try:
        with open(json_path, "w", encoding="utf-8") as jf:
            json.dump(record, jf, ensure_ascii=False, indent=4)
    except Exception as e:
        print(f"Lỗi khi lưu JSON: {e}")

    return True


def sync_excel_file() -> bool:
    """
    Tự động ghi và đồng bộ trực tiếp ra file Excel (ai_evaluation_results.xlsx)
    trên máy tính mỗi khi có thao tác lưu/cập nhật dữ liệu.
    """
    try:
        df = load_dataset(sort_by_case=True)
        excel_bytes = export_to_excel_bytes(df)
        with open(OUTPUT_EXCEL, "wb") as f:
            f.write(excel_bytes)
        return True
    except Exception as e:
        print(f"Lỗi khi đồng bộ file Excel: {e}")
        return False


def save_dataset(df: pd.DataFrame, create_backup: bool = True):
    """
    Lưu DataFrame trở lại file CSV chính, tự động tạo snapshot backup
    và tự động đồng bộ file Excel trên máy.
    """
    ensure_directories()
    if create_backup and os.path.exists(OUTPUT_CSV) and os.path.getsize(OUTPUT_CSV) > 0:
        create_snapshot_backup(note="auto_before_save")

    # Chuẩn hóa thứ tự cột
    save_df = df.copy()
    for col in CSV_HEADERS:
        if col not in save_df.columns:
            save_df[col] = ""
    save_df = save_df[CSV_HEADERS]

    save_df.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    sync_excel_file()


def append_evaluation_record(record: dict) -> bool:
    """
    Ghi thêm 1 bản ghi ca bệnh vào CSV, tạo file JSON metadata, tạo snapshot backup
    và tự động cập nhật file Excel trên máy.
    """
    ensure_directories()
    
    # 1. Tự động tạo snapshot backup
    create_snapshot_backup(note=f"case_{record.get('Case_ID', 'unknown')}")

    # 2. Ghi vào file CSV
    row = [record.get(col, "") for col in CSV_HEADERS]
    file_exists = os.path.exists(OUTPUT_CSV) and os.path.getsize(OUTPUT_CSV) > 0
    with open(OUTPUT_CSV, "a", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        if not file_exists:
            writer.writerow(CSV_HEADERS)
        writer.writerow(row)

    # 3. Lưu file JSON riêng biệt chứa metadata và toàn văn phản hồi AI
    case_id = record.get("Case_ID", "Case_000")
    safe_model = re.sub(r'[^a-zA-Z0-9_-]', '_', str(record.get("AI_Model", "model")))
    safe_prompt = re.sub(r'[^a-zA-Z0-9_-]', '_', str(record.get("Prompt_Type", "prompt")))
    ts_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    json_filename = f"{case_id}_{safe_model}_{safe_prompt}_{ts_str}.json"
    json_path = os.path.join(RESPONSES_DIR, json_filename)

    try:
        with open(json_path, "w", encoding="utf-8") as jf:
            json.dump(record, jf, ensure_ascii=False, indent=4)
    except Exception as e:
        print(f"Lỗi khi lưu JSON metadata: {e}")

    # 4. Tự động đồng bộ file Excel trên máy
    sync_excel_file()

    return True


def create_snapshot_backup(note: str = "") -> str:
    """
    Tạo bản sao lưu snapshot file CSV vào thư mục backups/
    """
    ensure_directories()
    if not os.path.exists(OUTPUT_CSV) or os.path.getsize(OUTPUT_CSV) == 0:
        return ""

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_note = re.sub(r'[^a-zA-Z0-9_-]', '_', note) if note else "manual"
    backup_filename = f"backup_{timestamp}_{safe_note}.csv"
    backup_path = os.path.join(BACKUPS_DIR, backup_filename)

    shutil.copy2(OUTPUT_CSV, backup_path)

    # Giới hạn giữ tối đa 30 bản sao lưu gần nhất
    clean_old_backups(max_keep=30)
    return backup_filename


def clean_old_backups(max_keep: int = 30):
    """Xóa các bản sao lưu cũ quá số lượng tối đa."""
    if not os.path.exists(BACKUPS_DIR):
        return
    backups = [
        os.path.join(BACKUPS_DIR, f) for f in os.listdir(BACKUPS_DIR)
        if f.startswith("backup_") and f.endswith(".csv")
    ]
    if len(backups) > max_keep:
        backups.sort(key=os.path.getmtime)
        for old_file in backups[:-max_keep]:
            try:
                os.remove(old_file)
            except OSError:
                pass


def list_backups() -> list:
    """
    Trả về danh sách thông tin các bản sao lưu có sẵn.
    """
    ensure_directories()
    if not os.path.exists(BACKUPS_DIR):
        return []

    files = [f for f in os.listdir(BACKUPS_DIR) if f.startswith("backup_") and f.endswith(".csv")]
    files.sort(key=lambda x: os.path.getmtime(os.path.join(BACKUPS_DIR, x)), reverse=True)

    backup_info = []
    for f in files:
        full_path = os.path.join(BACKUPS_DIR, f)
        mtime = datetime.fromtimestamp(os.path.getmtime(full_path)).strftime("%Y-%m-%d %H:%M:%S")
        size_kb = round(os.path.getsize(full_path) / 1024, 2)
        try:
            with open(full_path, "r", encoding="utf-8-sig", errors="ignore") as csv_f:
                row_count = sum(1 for _ in csv_f) - 1  # trừ header
                row_count = max(0, row_count)
        except Exception:
            row_count = "?"

        backup_info.append({
            "filename": f,
            "timestamp": mtime,
            "size_kb": size_kb,
            "row_count": row_count,
            "path": full_path
        })
    return backup_info


def restore_backup(backup_filename: str) -> bool:
    """
    Khôi phục dữ liệu từ một bản sao lưu snapshot.
    """
    ensure_directories()
    backup_path = os.path.join(BACKUPS_DIR, backup_filename)
    if not os.path.exists(backup_path):
        return False

    # Sao lưu trạng thái hiện tại trước khi khôi phục
    if os.path.exists(OUTPUT_CSV) and os.path.getsize(OUTPUT_CSV) > 0:
        create_snapshot_backup(note="pre_restore_safety")

    shutil.copy2(backup_path, OUTPUT_CSV)
    return True


def export_to_excel_bytes(df: pd.DataFrame) -> bytes:
    """
    Xuất dữ liệu thành file Excel (.xlsx) chuẩn nghiên cứu y khoa với 2 sheet:
    - Sheet 1: Dữ liệu Đánh giá (đầy đủ các dòng và format tiêu đề)
    - Sheet 2: Báo cáo Thống kê (Tổng quan chỉ số nghiên cứu)
    """
    wb = openpyxl.Workbook()
    # Sheet 1: Data
    ws_data = wb.active
    ws_data.title = "Dữ liệu Đánh giá"

    # Style tiêu đề
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # Ghi header
    cols = CSV_HEADERS
    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws_data.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # Ghi data
    for row_idx, (_, row) in enumerate(df.iterrows()):
        for col_idx, col_name in enumerate(cols, start=1):
            val = row.get(col_name, "")
            cell = ws_data.cell(row=row_idx + 2, column=col_idx, value=val)
            cell.border = border_thin
            if col_name in ["Time_Seconds", "Confidence"]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Tự động điều chỉnh độ rộng cột
    for col in ws_data.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > 50:
                val_str = val_str[:50]
            max_len = max(max_len, len(val_str))
        ws_data.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Sheet 2: Summary Report
    ws_summary = wb.create_sheet(title="Thống kê Nghiên cứu")
    ws_summary.column_dimensions['A'].width = 32
    ws_summary.column_dimensions['B'].width = 22

    summary_title_fill = PatternFill(start_color="203764", end_color="203764", fill_type="solid")
    title_cell = ws_summary.cell(row=1, column=1, value="BÁO CÁO TỔNG QUAN ĐÁNH GIÁ AI")
    title_cell.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    title_cell.fill = summary_title_fill
    ws_summary.merge_cells('A1:B1')

    # Tính toán chỉ số an toàn
    avg_time = "—"
    avg_conf = "—"
    if not df.empty:
        try:
            t_series = pd.to_numeric(df["Time_Seconds"], errors="coerce")
            avg_time = f"{round(t_series.mean(), 1)} giây"
        except Exception:
            pass
        try:
            c_series = pd.to_numeric(df["Confidence"], errors="coerce")
            avg_conf = f"{round(c_series.mean(), 1)}%"
        except Exception:
            pass

    metrics = [
        ("Thời gian xuất báo cáo", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Tổng số ca đã đánh giá", len(df)),
        ("Thời gian trung bình mỗi ca", avg_time),
        ("Độ tin cậy AI trung bình", avg_conf),
    ]

    r = 3
    for label, val in metrics:
        ws_summary.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws_summary.cell(row=r, column=2, value=val)
        r += 1

    r += 1
    ws_summary.cell(row=r, column=1, value="PHÂN BỔ THEO MÔ HÌNH AI").font = Font(bold=True, color="1F4E79")
    r += 1
    if not df.empty and "AI_Model" in df.columns:
        model_counts = df["AI_Model"].value_counts()
        for m_name, count in model_counts.items():
            ws_summary.cell(row=r, column=1, value=str(m_name))
            ws_summary.cell(row=r, column=2, value=int(count))
            r += 1

    r += 1
    ws_summary.cell(row=r, column=1, value="PHÂN BỔ PELL & GREGORY").font = Font(bold=True, color="1F4E79")
    r += 1
    if not df.empty and "Pell_Gregory_Full" in df.columns:
        pg_counts = df["Pell_Gregory_Full"].value_counts()
        for pg_name, count in pg_counts.items():
            ws_summary.cell(row=r, column=1, value=str(pg_name))
            ws_summary.cell(row=r, column=2, value=int(count))
            r += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def create_ai_responses_zip() -> bytes:
    """
    Nén toàn bộ các file JSON trong thư mục ai_responses thành file ZIP bytes.
    """
    ensure_directories()
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        if os.path.exists(RESPONSES_DIR):
            for root, _, files in os.walk(RESPONSES_DIR):
                for file in files:
                    file_path = os.path.join(root, file)
                    arcname = os.path.relpath(file_path, RESPONSES_DIR)
                    zf.write(file_path, arcname=arcname)
    buf.seek(0)
    return buf.getvalue()


def load_session_progress() -> dict:
    """
    Tải tiến độ phiên làm việc từ session_progress.json.
    """
    if os.path.exists(PROGRESS_JSON):
        try:
            with open(PROGRESS_JSON, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass

    # Mặc định danh sách 129 ca (Case_001 -> Case_129)
    cases = list(ALL_CASES)
    return {
        "cases_remaining": cases,
        "cases_done": [],
        "operator": OPERATORS_LIST[0],
        "model": MODELS_LIST[0],
        "prompt": "P2 – Structured Medical"
    }


def save_session_progress(progress_data: dict):
    """
    Lưu tiến độ phiên làm việc vào session_progress.json.
    """
    try:
        with open(PROGRESS_JSON, "w", encoding="utf-8") as f:
            json.dump(progress_data, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"Lỗi khi lưu tiến độ phiên: {e}")


def parse_ai_response(raw_text: str) -> dict:
    """
    Tự động phân tích và trích xuất các trường phân loại y khoa từ phản hồi AI
    (hỗ trợ cả định dạng JSON lẫn văn bản tự nhiên, Markdown, Chain-of-Thought).
    """
    if not raw_text or len(raw_text.strip()) < 4:
        return {}

    text = raw_text.replace('“', '"').replace('”', '"').replace('‘', "'").replace('’', "'")
    result = {}

    # Helper: Chuẩn hóa phân loại Class sang "I", "II", "III"
    def normalize_class(val_str: str) -> str:
        if not val_str:
            return ""
        val_str = str(val_str).strip().upper()
        m = re.search(r'\b(III|II|I|3|2|1)\b', val_str, re.I)
        if m:
            v = m.group(1).upper()
            if v in ("1", "I"): return "I"
            if v in ("2", "II"): return "II"
            if v in ("3", "III"): return "III"
        return ""

    # Helper: Chuẩn hóa phân loại Position sang "A", "B", "C"
    def normalize_position(val_str: str) -> str:
        if not val_str:
            return ""
        val_str = str(val_str).strip().upper()
        m = re.search(r'\b([ABC])\b', val_str)
        if m:
            return m.group(1).upper()
        return ""

    # Helper: Chuẩn hóa Winter's Classification
    def normalize_winter(val_str: str) -> str:
        if not val_str:
            return ""
        s = str(val_str).strip()
        m = re.search(r'\b(mesio[- ]?angular|disto[- ]?angular|horizontal|vertical|bucco[- ]?lingual|buccal|lingual|others?|transverse|inverted)\b', s, re.I)
        if m:
            w = m.group(1).lower().replace("-", "").replace(" ", "")
            if "mesio" in w: return "Mesioangular"
            if "disto" in w: return "Distoangular"
            if "horiz" in w: return "Horizontal"
            if "vert" in w: return "Vertical"
            if "bucco" in w or "buccal" in w or "lingual" in w: return "Buccolingual"
            return "Others"
        return ""

    # Helper: Chuẩn hóa Confidence
    def normalize_conf(val) -> int:
        try:
            if isinstance(val, (int, float)):
                v = int(val)
                if 0 <= v <= 1: v = int(v * 100)
                if 0 <= v <= 100: return v
            s = str(val).strip()
            m = re.search(r'(\d{1,3})\s*%', s)
            if not m:
                m = re.search(r'\b([1-9]\d?|100)\b', s)
            if m:
                v = int(m.group(1))
                if 0 <= v <= 100: return v
        except Exception:
            pass
        return None

    # Helper: Chuẩn hóa Pederson Level
    def normalize_pederson(val_str: str) -> str:
        s = str(val_str)
        if re.search(r'\b(nhẹ|mild|easy|3[-–]4)\b', s, re.I):
            return "Nhẹ (3–4)"
        if re.search(r'\b(trung bình|moderate|medium|5[-–]6)\b', s, re.I):
            return "Trung bình (5–6)"
        if re.search(r'\b(khó|difficult|high|hard|7[-–]10)\b', s, re.I):
            return "Khó (7–10)"
        return ""

    # ─────────────────────────────────────────────────────────────
    # GIAI ĐOẠN 1: Bóc tách khối JSON (Code blocks hoặc Raw JSON)
    # ─────────────────────────────────────────────────────────────
    json_candidates = []
    for match in re.finditer(r'```(?:json)?\s*(\{[\s\S]*?\})\s*```', text, re.I):
        json_candidates.append(match.group(1))
    for match in re.finditer(r'\{[^{}]*(?:p&g|winter|class|confidence)[^{}]*\}', text, re.I):
        json_candidates.append(match.group(0))

    for j_str in json_candidates:
        clean_j = re.sub(r'//[^\n\r]*', '', j_str)
        clean_j = re.sub(r',\s*\}', '}', clean_j)
        try:
            data = json.loads(clean_j)
            if isinstance(data, dict):
                for k, v in data.items():
                    k_clean = str(k).lower().replace("_", "").replace(" ", "").replace("-", "").replace("&", "")
                    v_str = str(v)

                    # Class
                    if "class" in k_clean and "winter" not in k_clean and "position" not in k_clean:
                        cls = normalize_class(v_str)
                        if cls and "Pell_Gregory_Class" not in result:
                            result["Pell_Gregory_Class"] = cls
                    
                    # Position
                    elif "position" in k_clean or k_clean in ("pos", "pgpos", "pgposition"):
                        pos = normalize_position(v_str)
                        if pos and "Pell_Gregory_Position" not in result:
                            result["Pell_Gregory_Position"] = pos

                    # Combined P&G field
                    elif ("pell" in k_clean or "pg" in k_clean) and "winter" not in k_clean:
                        m_cls = normalize_class(v_str)
                        m_pos = normalize_position(v_str)
                        if m_cls and "Pell_Gregory_Class" not in result:
                            result["Pell_Gregory_Class"] = m_cls
                        if m_pos and "Pell_Gregory_Position" not in result:
                            result["Pell_Gregory_Position"] = m_pos

                    # Winter
                    elif "winter" in k_clean or "angulation" in k_clean:
                        w = normalize_winter(v_str)
                        if w and "Winter_Class" not in result:
                            result["Winter_Class"] = w

                    # Confidence
                    elif "conf" in k_clean:
                        c = normalize_conf(v)
                        if c is not None and "Confidence" not in result:
                            result["Confidence"] = c

                    # Pederson
                    elif "pederson" in k_clean or "difficulty" in k_clean:
                        ped = normalize_pederson(v_str)
                        if ped and "Pederson_Level" not in result:
                            result["Pederson_Level"] = ped
        except Exception:
            pass

    if "Pell_Gregory_Class" in result and "Pell_Gregory_Position" in result and "Winter_Class" in result:
        return result

    # ─────────────────────────────────────────────────────────────
    # GIAI ĐOẠN 2: Bóc tách văn bản tự nhiên (Text Parsing)
    # ─────────────────────────────────────────────────────────────
    # 1. Tìm phần Summary / Conclusion nếu có
    summary_text = ""
    summary_match = re.search(
        r'(?:###?\s*)?(?:Summary|Conclusion|Final Diagnosis|Final Classification|Impression|Kết luận|Tóm tắt|Results?|Overall)[:\s\n]+([\s\S]+)$',
        text,
        re.I
    )
    if summary_match:
        summary_text = summary_match.group(1).strip()

    # 2. Xóa các dòng định nghĩa lý thuyết từ prompt để tránh bắt nhầm Class I / Pos A
    cleaned_full_text = re.sub(r'Class\s*(?:I|II|III|1|2|3)\s*:\s*Crown\s+[^.\n\r]+[.\n\r]?', '', text, flags=re.I)
    cleaned_full_text = re.sub(r'Position\s*[ABC]\s*:\s*Crown\s+[^.\n\r]+[.\n\r]?', '', cleaned_full_text, flags=re.I)

    scan_scopes = []
    if summary_text:
        scan_scopes.append(summary_text)
    scan_scopes.append(cleaned_full_text)
    scan_scopes.append(text)

    # ── A. Trích xuất Pell & Gregory ───────────────────────────
    for scope in scan_scopes:
        if "Pell_Gregory_Class" in result and "Pell_Gregory_Position" in result:
            break

        # Bỏ phần chẩn đoán phụ trong ngoặc đơn "(or ...)" để ưu tiên chẩn đoán chính
        scope_clean = re.sub(r'\(or\s+[^)]+\)', '', scope, flags=re.I)
        
        # A1. Dạng kết hợp trực tiếp: "Class II, Position B" hoặc "P&G: Class II, Position B"
        combo_match = re.search(
            r'(?:Pell\s*(?:&|and)?\s*Gregory|P&G)?[\s\w:\-]*?Class\s*(III|II|I|3|2|1)[,\s/]+(?:Position\s*|Pos\s*)?([ABC])\b',
            scope_clean,
            re.I
        )
        if combo_match:
            if "Pell_Gregory_Class" not in result:
                result["Pell_Gregory_Class"] = normalize_class(combo_match.group(1))
            if "Pell_Gregory_Position" not in result:
                result["Pell_Gregory_Position"] = normalize_position(combo_match.group(2))
            break

        # A2. Dạng mã ngắn: "II-B", "II/B", "II - B", "Class II-B"
        code_match = re.search(r'\b(?:Class\s*)?(III|II|I)\s*[-–/]\s*([A-C])\b', scope_clean, re.I)
        if code_match:
            if "Pell_Gregory_Class" not in result:
                result["Pell_Gregory_Class"] = normalize_class(code_match.group(1))
            if "Pell_Gregory_Position" not in result:
                result["Pell_Gregory_Position"] = normalize_position(code_match.group(2))
            break

        # A3. Từng trường riêng lẻ có nhãn rõ ràng
        if "Pell_Gregory_Class" not in result:
            c_m = re.search(r'(?:Pell\s*(?:&|and)?\s*Gregory|P&G|P&G\s*Class|Class)\s*[:=]\s*(?:Class\s*)?(III|II|I|3|2|1)\b', scope_clean, re.I)
            if c_m:
                result["Pell_Gregory_Class"] = normalize_class(c_m.group(1))

        if "Pell_Gregory_Position" not in result:
            p_m = re.search(r'(?:Position|Pos|P&G\s*Position)\s*[:=]\s*(?:Position\s*)?([ABC])\b', scope_clean, re.I)
            if p_m:
                result["Pell_Gregory_Position"] = normalize_position(p_m.group(1))

    # ── B. Trích xuất Winter's Classification ──────────────────
    if "Winter_Class" not in result:
        for scope in scan_scopes:
            w_labeled = re.search(r'(?:Winter(?:[\'’]s)?|Angulation|Thế răng)[^\n\r:=]*[:=]\s*([^\n\r,;.]+)', scope, re.I)
            if w_labeled:
                w_val = normalize_winter(w_labeled.group(1))
                if w_val:
                    result["Winter_Class"] = w_val
                    break

            w_direct = normalize_winter(scope)
            if w_direct:
                result["Winter_Class"] = w_direct
                break

    # ── C. Trích xuất Confidence (Độ tin cậy) ─────────────────
    if "Confidence" not in result:
        for scope in scan_scopes:
            c_m = re.search(r'(?:Confidence|Độ tin cậy|Confidence level)[^\n\r:=]*[:=]\s*(\d{1,3})\s*%?', scope, re.I)
            if c_m:
                val = int(c_m.group(1))
                if 0 <= val <= 100:
                    result["Confidence"] = val
                    break

            c_rev = re.search(r'(\d{1,3})\s*%\s*(?:confidence|độ tin cậy)?', scope, re.I)
            if c_rev:
                val = int(c_rev.group(1))
                if 0 <= val <= 100:
                    result["Confidence"] = val
                    break

    # ── D. Trích xuất Pederson Level ─────────────────────────
    if "Pederson_Level" not in result:
        for scope in scan_scopes:
            ped = normalize_pederson(scope)
            if ped:
                result["Pederson_Level"] = ped
                break

    return result
